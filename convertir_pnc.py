# -*- coding: utf-8 -*-
"""
convertir_pnc.py — Lee el/los Excel de novedades (PNC) y genera data_pnc.js
con el resumen de PNC por placa, para cruzarlo con la app de viajes.

Por cada placa calcula:
  - PNC totales, abiertos y cerrados
  - PNC atribuibles al transportador (imputación TRANSPORTADOR o cobro SI ATRIBUIBLE)
  - Desglose por TIPO DE PNC (devolución, averiado, faltante, etc.)
  - Total descontado al afiliado (Dscto Afiliado)
  - Fecha del último PNC
  - Semáforo de riesgo según % de PNC atribuibles sobre el total

USO:
  1. Edita RUTA_CARPETA_PNC con tu carpeta de OneDrive (Analisis PNC).
     (También sirve si apuntas a un solo archivo .xlsx)
  2. Ejecuta:  python convertir_pnc.py   (o actualizar.bat que corre todo)

Requiere:  pip install openpyxl
"""

import json
import os
import sys
import glob
from datetime import datetime, date

# ============================================================
# ⚙️ CONFIGURACIÓN — EDITA ESTA LÍNEA CON TU CARPETA REAL
# Carpeta que contiene el/los Excel de PNC. Ejemplo:
#   C:/Users/jarias/OneDrive - TRACTOCAR LOGISTICS SAS/Cartera_Afiliados/Analisis PNC
# ============================================================
RUTA_CARPETA_PNC = "C:/Users/jarias/OneDrive - TRACTOCAR LOGISTICS SAS/Cartera_Afiliados/Analisis PNC/Pnc Analizado 3.xlsx"

# Nombre de la hoja con los datos (si no la encuentra, busca una que tenga las columnas clave)
HOJA_PNC = "DATOS PNC"

SALIDA = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data_pnc.js")

# Archivo con los manifiestos válidos (lo genera convertir_viajes.py).
# El PNC solo cuenta novedades cuyo MANIFIESTO exista en los viajes cargados,
# así PNC y viajes hablan del mismo periodo (2025 en adelante).
RUTA_MANIFIESTOS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "manifiestos_validos.json")

# Umbral del semáforo: % de PNC atribuibles sobre el total de la placa
PCT_RIESGO_ALTO = 40   # >= 40% atribuibles -> riesgo ALTO
PCT_RIESGO_MEDIO = 15  # 15%-40% -> MEDIO ; por debajo -> BAJO


def limpio(v):
    if v is None:
        return ""
    s = str(v).strip()
    return "" if s.lower() in {"-", "n/d", "nan", "nat", "none"} else s


def a_fecha(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = limpio(v)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s[:19], fmt).date()
        except ValueError:
            pass
    return None


def a_numero(v):
    if v is None:
        return 0.0
    try:
        return float(v)
    except (ValueError, TypeError):
        s = limpio(v).replace("$", "").replace(".", "").replace(",", ".")
        try:
            return float(s)
        except ValueError:
            return 0.0


def norm_tipo(t):
    """Normaliza el TIPO DE PNC para que 'DEVOLUCION ' y 'DEVOLUCION' sean iguales."""
    t = limpio(t).upper().strip()
    if not t:
        return "SIN CLASIFICAR"
    return t


def norm_manifiesto(v):
    """Quita TCL. y .0 para que el manifiesto del PNC cruce con el Envio de viajes."""
    s = limpio(v).replace("TCL.", "").strip()
    if s.endswith(".0"):
        s = s[:-2]
    return s


def es_atribuible(imputacion, cobro):
    imp = limpio(imputacion).upper()
    cob = limpio(cobro).upper()
    if imp == "TRANSPORTADOR":
        return True
    if cob.startswith("SI ATRIBUIBLE"):
        return True
    return False


def main():
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("❌ Falta openpyxl. Instálala con:  pip install openpyxl")
        sys.exit(1)

    # Cargar manifiestos válidos (del periodo de viajes). Si no existe el archivo,
    # se cuentan todos los PNC y se avisa.
    manifiestos_validos = None
    if os.path.exists(RUTA_MANIFIESTOS):
        try:
            with open(RUTA_MANIFIESTOS, encoding="utf-8") as f:
                manifiestos_validos = set(json.load(f))
            print(f"🔗 Cruzando con {len(manifiestos_validos):,} manifiestos de viajes (solo PNC de ese periodo).")
        except Exception as e:
            print(f"⚠️ No pude leer manifiestos_validos.json ({e}); contaré todos los PNC.")
    else:
        print("⚠️ No existe manifiestos_validos.json. Ejecuta primero convertir_viajes.py")
        print("   Por ahora cuento TODOS los PNC (incluye años anteriores a 2025).")

    # Aceptar carpeta o archivo único
    if os.path.isfile(RUTA_CARPETA_PNC):
        archivos = [RUTA_CARPETA_PNC]
    elif os.path.isdir(RUTA_CARPETA_PNC):
        archivos = sorted(
            f for f in glob.glob(os.path.join(RUTA_CARPETA_PNC, "*.xls*"))
            if not os.path.basename(f).startswith("~$")
        )
    else:
        print("❌ No encuentro la carpeta (o archivo) de PNC:")
        print("   " + RUTA_CARPETA_PNC)
        print("   Edita RUTA_CARPETA_PNC al inicio de convertir_pnc.py")
        sys.exit(1)

    if not archivos:
        print("❌ La carpeta no tiene archivos .xlsx:")
        print("   " + RUTA_CARPETA_PNC)
        sys.exit(1)

    print(f"📂 PNC: {RUTA_CARPETA_PNC}")
    print(f"   {len(archivos)} archivo(s)\n")

    COLS = ["PLACA VEHICULO", "FECHA DE REGISTRO", "Estado del RPNC",
            "TIPO DE PNC", "IMPUTACION", "COBRO PROCEDENTE", "Dscto Afiliado",
            "MANIFIESTO", "NO CONFORMIDAD"]

    flota = {}   # placa -> acumulador
    filas_totales = 0
    fmin, fmax = None, None

    for ruta in archivos:
        nombre = os.path.basename(ruta)
        try:
            wb = load_workbook(ruta, read_only=True, data_only=True)
        except Exception as e:
            print(f"   ⚠️ {nombre}: no se pudo abrir ({e}) — lo salto")
            continue

        # localizar hoja
        hoja, encabezados = None, None
        if HOJA_PNC in wb.sheetnames:
            primera = next(wb[HOJA_PNC].iter_rows(max_row=1, values_only=True), None)
            if primera and "PLACA VEHICULO" in primera:
                hoja, encabezados = HOJA_PNC, list(primera)
        if not hoja:
            for nh in wb.sheetnames:
                primera = next(wb[nh].iter_rows(max_row=1, values_only=True), None)
                if primera and "PLACA VEHICULO" in primera and "Estado del RPNC" in primera:
                    hoja, encabezados = nh, list(primera)
                    break
        if not hoja:
            print(f"   ⚠️ {nombre}: no encontré la hoja de PNC — lo salto")
            continue

        idx = {n: i for i, n in enumerate(encabezados) if n}

        def col(fila, nombre_col):
            i = idx.get(nombre_col)
            return fila[i] if i is not None and i < len(fila) else None

        n = 0
        omitidos = 0
        for fila in wb[hoja].iter_rows(min_row=2, values_only=True):
            placa = limpio(col(fila, "PLACA VEHICULO")).upper().replace("TCL.", "")
            if not placa:
                continue

            # Filtro por manifiesto: solo PNC cuyo manifiesto esté en los viajes cargados
            manif = norm_manifiesto(col(fila, "MANIFIESTO"))
            if manifiestos_validos is not None:
                if not manif or manif not in manifiestos_validos:
                    omitidos += 1
                    continue

            n += 1
            filas_totales += 1

            estado = limpio(col(fila, "Estado del RPNC")).upper()
            tipo = norm_tipo(col(fila, "TIPO DE PNC"))
            fecha = a_fecha(col(fila, "FECHA DE REGISTRO"))
            atrib = es_atribuible(col(fila, "IMPUTACION"), col(fila, "COBRO PROCEDENTE"))
            dscto = a_numero(col(fila, "Dscto Afiliado"))

            if fecha:
                fmin = fecha if not fmin or fecha < fmin else fmin
                fmax = fecha if not fmax or fecha > fmax else fmax

            f = flota.setdefault(placa, {
                "total": 0, "abiertos": 0, "cerrados": 0,
                "atribuibles": 0, "atrib_abiertos": 0,
                "dscto": 0.0, "tipos": {}, "ult": None,
            })
            f["total"] += 1
            if estado == "ABIERTO":
                f["abiertos"] += 1
            elif estado == "CERRADO":
                f["cerrados"] += 1
            if atrib:
                f["atribuibles"] += 1
                if estado == "ABIERTO":
                    f["atrib_abiertos"] += 1
            f["dscto"] += dscto
            f["tipos"][tipo] = f["tipos"].get(tipo, 0) + 1
            if fecha and (not f["ult"] or fecha > f["ult"]):
                f["ult"] = fecha

        if manifiestos_validos is not None:
            print(f"   ✅ {nombre} ({hoja}): {n:,} PNC del periodo · {omitidos:,} omitidos (sin viaje en 2025+)")
        else:
            print(f"   ✅ {nombre} ({hoja}): {n:,} filas")

    if not flota:
        print("❌ No se pudo leer ningún dato de PNC.")
        sys.exit(1)

    # consolidar
    placas_out = []
    for placa, f in flota.items():
        total = f["total"]
        atrib = f["atribuibles"]
        pct = (atrib / total * 100) if total else 0
        if pct >= PCT_RIESGO_ALTO and atrib >= 2:
            riesgo = "ALTO"
        elif pct >= PCT_RIESGO_MEDIO and atrib >= 1:
            riesgo = "MEDIO"
        else:
            riesgo = "BAJO"

        tipos = sorted(
            [{"t": k, "n": v} for k, v in f["tipos"].items()],
            key=lambda x: -x["n"]
        )
        tipo_top = tipos[0]["t"] if tipos else ""

        placas_out.append({
            "placa": placa,
            "total": total,
            "abiertos": f["abiertos"],
            "cerrados": f["cerrados"],
            "atribuibles": atrib,
            "atribAbiertos": f["atrib_abiertos"],
            "pctAtrib": round(pct),
            "riesgo": riesgo,
            "dscto": round(f["dscto"]),
            "tipoTop": tipo_top,
            "tipos": tipos,
            "ultimo": f["ult"].strftime("%Y-%m-%d") if f["ult"] else "",
        })

    placas_out.sort(key=lambda x: -x["atribuibles"])

    paquete = {
        "actualizado": datetime.now().strftime("%d/%m/%Y %H:%M"),
        "desde": fmin.strftime("%Y-%m-%d") if fmin else "",
        "hasta": fmax.strftime("%Y-%m-%d") if fmax else "",
        "archivos": len(archivos),
        "filas": filas_totales,
        "placas": placas_out,
    }

    with open(SALIDA, "w", encoding="utf-8") as fp:
        fp.write("// Generado por convertir_pnc.py — NO editar a mano.\n")
        fp.write("window.DATOS_PNC = ")
        json.dump(paquete, fp, ensure_ascii=False, separators=(",", ":"))
        fp.write(";\n")

    peso = os.path.getsize(SALIDA) / 1024 / 1024
    abiertos = sum(p["abiertos"] for p in placas_out)
    atrib = sum(p["atribuibles"] for p in placas_out)
    alto = sum(1 for p in placas_out if p["riesgo"] == "ALTO")
    print(f"\n✅ data_pnc.js generado — {peso:.1f} MB")
    print(f"   {len(placas_out):,} placas con PNC · {filas_totales:,} PNC · {fmin} → {fmax}")
    print(f"   PNC abiertos: {abiertos:,} · atribuibles al transportador: {atrib:,}")
    print(f"   Placas en riesgo ALTO: {alto:,}")


if __name__ == "__main__":
    main()
