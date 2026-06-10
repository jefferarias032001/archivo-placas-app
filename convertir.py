# -*- coding: utf-8 -*-
"""
convertir.py — Convierte ARCHIVO DE PLACAS.xlsx en data.js para la web app.

USO:
  1. Edita la variable RUTA_EXCEL aquí abajo con la ruta de tu archivo
     en la carpeta sincronizada de OneDrive.
  2. Ejecuta:  python convertir.py
  3. Se genera (o actualiza) el archivo data.js en esta misma carpeta.

Requiere una sola librería:  pip install openpyxl
"""

import json
import os
import sys
from datetime import datetime, date

# ============================================================
# ⚙️ CONFIGURACIÓN — EDITA ESTA LÍNEA CON TU RUTA REAL
# Ejemplos típicos de ruta de OneDrive en Windows:
#   C:/Users/TuUsuario/OneDrive - NombreEmpresa/ARCHIVO DE PLACAS.xlsx
#   C:/Users/TuUsuario/OneDrive/Documentos/ARCHIVO DE PLACAS.xlsx
# (usa barras / o dobles \\ — no barras simples \)
# ============================================================
RUTA_EXCEL = "C:/Users/TU_USUARIO/OneDrive - TU_EMPRESA/ARCHIVO DE PLACAS.xlsx"

SALIDA = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data.js")

VACIOS = {"-", "- -", "N/D", "0", "0000000", "", None}


def limpio(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return "" if s in VACIOS else s


def norm_tipo_base(t):
    if not t:
        return "SIN TIPOLOGÍA"
    base = t.split("_")[0].strip()
    if base.endswith("REFRIGERADA"):
        base = base[: -len("REFRIGERADA")].strip()
    base = base.replace("DOBLE TROQUE", "DOBLETROQUE").replace("DOBLE-TROQUE", "DOBLETROQUE")
    if base == "PATINTA":
        base = "PATINETA"
    return base or "SIN TIPOLOGÍA"


def a_fecha(v):
    if isinstance(v, datetime):
        return v
    if isinstance(v, date):
        return datetime(v.year, v.month, v.day)
    if isinstance(v, str):
        for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(v.strip()[:10], fmt)
            except ValueError:
                pass
    return datetime(1900, 1, 1)


def main():
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("❌ Falta la librería openpyxl. Instálala con:  pip install openpyxl")
        sys.exit(1)

    if not os.path.exists(RUTA_EXCEL):
        print("❌ No encuentro el archivo de Excel en:")
        print("   " + RUTA_EXCEL)
        print("   Edita la variable RUTA_EXCEL al inicio de convertir.py")
        sys.exit(1)

    print("📖 Leyendo:", RUTA_EXCEL)
    wb = load_workbook(RUTA_EXCEL, read_only=True, data_only=True)

    # Detectar la hoja que tenga las columnas PLACA y TIPOLOGIA
    hoja, encabezados = None, None
    for nombre in wb.sheetnames:
        ws = wb[nombre]
        primera = next(ws.iter_rows(max_row=1, values_only=True), None)
        if primera and "PLACA" in primera and "TIPOLOGIA" in primera:
            hoja, encabezados = nombre, list(primera)
            break
    if not hoja:
        print("❌ Ninguna hoja tiene las columnas PLACA y TIPOLOGIA.")
        sys.exit(1)
    print("✅ Hoja detectada:", hoja)

    idx = {nombre: i for i, nombre in enumerate(encabezados) if nombre}

    def col(fila, nombre):
        i = idx.get(nombre)
        return limpio(fila[i]) if i is not None and i < len(fila) else ""

    por_placa = {}
    total = 0
    for fila in wb[hoja].iter_rows(min_row=2, values_only=True):
        placa = col(fila, "PLACA")
        if not placa:
            continue
        total += 1
        i_fecha = idx.get("FECHA DE CREACIÓN")
        fecha = a_fecha(fila[i_fecha]) if i_fecha is not None and i_fecha < len(fila) else datetime(1900, 1, 1)
        previo = por_placa.get(placa)
        if previo and previo["_fecha"] >= fecha:
            continue
        tip = col(fila, "TIPOLOGIA") or "-"
        por_placa[placa] = {
            "_fecha": fecha,
            "placa": placa,
            "estado": col(fila, "ESTADO").upper(),
            "tipologia": tip,
            "tipoBase": norm_tipo_base(col(fila, "TIPOLOGIA")),
            "carroceria": col(fila, "CARROCERÍA"),
            "marca": col(fila, "MARCA"),
            "modelo": col(fila, "MODELO"),
            "color": col(fila, "COLOR"),
            "linea": col(fila, "LÍNEA"),
            "remolque": col(fila, "REMOLQUE"),
            "cond": {
                "nombre": col(fila, "NOMBRE Y APELLIDO"),
                "cedula": col(fila, "CEDULA"),
                "celular": col(fila, "CELULAR"),
                "ciudad": col(fila, "CIUDAD"),
                "direccion": col(fila, "DIRECCIÓN"),
                "licencia": col(fila, "LICENCIA"),
                "clase": col(fila, "CLASE"),
                "vence": col(fila, "FECHA DE VENCIMIENTO"),
            },
            "tene": {
                "nombre": col(fila, "NOMBRE_2"),
                "cedula": col(fila, "CEDULA_1"),
                "celular": col(fila, "CELULAR_3"),
                "ciudad": col(fila, "CIUDAD_5"),
                "direccion": col(fila, "DIRECCIÓN_4"),
            },
            "prop": {
                "nombre": col(fila, "NOMBRE_7"),
                "cedula": col(fila, "CEDULA_6"),
                "celular": col(fila, "CELULAR_8"),
                "ciudad": col(fila, "CIUDAD_9"),
                "direccion": col(fila, "DIRECCIÓN_10"),
            },
        }

    registros = sorted(por_placa.values(), key=lambda d: d["placa"])
    for r in registros:
        del r["_fecha"]

    paquete = {
        "actualizado": datetime.now().strftime("%d/%m/%Y %H:%M"),
        "totalRegistros": total,
        "hoja": hoja,
        "placas": registros,
    }

    with open(SALIDA, "w", encoding="utf-8") as f:
        f.write("// Archivo generado automáticamente por convertir.py — NO editar a mano.\n")
        f.write("// Contiene datos personales: este repositorio debe ser PRIVADO.\n")
        f.write("window.DATOS_PLACAS = ")
        json.dump(paquete, f, ensure_ascii=False, separators=(",", ":"))
        f.write(";\n")

    peso = os.path.getsize(SALIDA) / 1024 / 1024
    print(f"✅ data.js generado: {len(registros):,} placas únicas (de {total:,} registros) — {peso:.1f} MB")
    print("   Siguiente paso: subir a GitHub (o usa actualizar.bat que hace todo junto).")


if __name__ == "__main__":
    main()
